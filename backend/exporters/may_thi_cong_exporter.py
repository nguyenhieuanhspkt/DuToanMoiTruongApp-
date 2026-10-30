import os
import json
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any, List, Tuple

def build_may_thi_cong_sheet(
    wb, 
    last_hp_row: int,
    jobs_dtos: List[Any],
    font_title_15b: Font,
    font_sub_11i: Font,
    font_header_11b: Font,
    font_data_11r: Font,
    font_data_11b: Font,
    fill_header_gray: PatternFill,
    fill_sub_gray: PatternFill,
    fill_total_pale: PatternFill,
    border_cell: Border,
    border_total: Border,
    FMT_CURRENCY: str,
    FMT_NORM: str
) -> Tuple[Worksheet, int]:
    """
    Xây dựng sheet 'Máy thi công' chuẩn mẫu F1:
    - Lọc bỏ triệt để các dòng tiêu đề phân mục dựa trên STT dạng số.
    - Cột B dòng con: Link công thức động trực tiếp sang STT của sheet 'Công trình' (='Công trình'!A{row}).
    - Cột E dòng con: Link tên công tác trực tiếp từ sheet 'Công trình' (='Công trình'!D{row}).
    - Dùng SUMIFS quét chính xác tổng ca máy từ sheet 'Hao phí VT' theo đúng Mã công tác và Tên thiết bị.
    - Bung dòng con truy vết công tác màu xanh (có Group ẩn/hiện) và đánh số STT chuẩn từ 1 đến n.
    """
    if "Máy thi công" in wb.sheetnames:
        ws_m = wb["Máy thi công"]
        wb.remove(ws_m)
    
    ws_m = wb.create_sheet(title="Máy thi công")
    ws_m.views.sheetView[0].showGridLines = True

    # 1. KHỐI TIÊU ĐỀ CHUẨN F1
    ws_m.append(["BẢNG TỔNG HỢP MÁY THI CÔNG VÀ CHÊNH LỆCH GIÁ"])
    ws_m.append(["(Căn cứ theo Thông tư 20/2017/TT-BTNMT và định giá ca máy hiện hành)"])
    ws_m.append(["='Công trình'!A3"])  # Link tên công trình động
    ws_m.append(["='Công trình'!A4"])  # Link tên hạng mục động
    ws_m.append([])
    
    # Dòng đơn vị tính đặt ở góc trên phải (cột Q - index 16)
    ws_m.cell(5, 17, "Đơn vị: đồng")

    # Header bảng (Dòng 6) - Tổng cộng 17 cột (A đến Q) khớp mẫu chuẩn F1
    headers = [""] * 17
    headers[0] = "STT"           # Cột A
    headers[1] = "Mã hiệu"       # Cột B
    headers[4] = "Tên vật tư"    # Cột E (Hiển thị tên máy thi công)
    headers[7] = "Đơn vị"        # Cột H
    headers[9] = "Khối lượng"    # Cột J
    headers[10] = "Giá gốc"      # Cột K
    headers[11] = "Giá thông báo"# Cột L
    headers[12] = "Bù giá"       # Cột M
    headers[13] = "Giá hiện tại" # Cột N
    headers[14] = "Chênh lệch"   # Cột O
    headers[15] = "Tổng chênh"   # Cột P
    headers[16] = "Thành tiền"   # Cột Q
    ws_m.append(headers)

    # 2. TỰ ĐỘNG BÓC TÁCH DANH MỤC MÁY MÓC/THIẾT BỊ TỪ JOBS_DTOS (Loại bỏ dòng tiêu đề phân mục)
    equip_map = {}
    for job in jobs_dtos:
        # Kiểm tra chuẩn: Chỉ lấy công tác có STT dạng số hợp lệ (bỏ qua dòng tiêu đề phân mục như 'PHẦN: ...')
        raw_stt = getattr(job, 'stt', None)
        if not raw_stt or not str(raw_stt).strip().isdigit():
            continue

        m_group = job.groups.get("M")
        if m_group and hasattr(m_group, 'items'):
            for detail_item in m_group.items:
                eq_name = detail_item.name.strip()
                eq_unit = detail_item.unit.strip() if detail_item.unit else "ca"
                
                # Gom nhóm các thiết bị trùng tên để tạo danh mục cha độc lập duy nhất
                if eq_name not in equip_map:
                    equip_map[eq_name] = {
                        "id": f"M_{len(equip_map)+1:03d}",
                        "name": eq_name,
                        "unit": eq_unit,
                        "price": 0.0 
                    }
    
    equip_items = list(equip_map.values())

    m_start_row = 7
    curr_row = m_start_row

    # 3. VÒNG LẶP DUYỆT QUA TỪNG MÁY THI CÔNG CHÍNH (DÒNG CHA)
    for idx, it in enumerate(equip_items, 1):
        m_id = it.get('id')
        m_name = it.get('name')
        m_unit = it.get('unit')
        price = float(it.get('price', 0.0))

        # Thu thập các công tác sử dụng máy này từ jobs_dtos để làm dòng con truy vết
        usage_jobs = []
        for job in jobs_dtos:
            # Lọc chặn các dòng tiêu đề phân mục không có STT số
            raw_stt = getattr(job, 'stt', None)
            if not raw_stt or not str(raw_stt).strip().isdigit():
                continue

            m_group = job.groups.get("M")
            if m_group and hasattr(m_group, 'items'):
                for detail_item in m_group.items:
                    detail_name_clean = detail_item.name.strip().lower()
                    m_name_clean = m_name.strip().lower()
                    
                    if detail_name_clean in m_name_clean or m_name_clean in detail_name_clean:
                        # [CƠ CHẾ LINK KHỐI LƯỢNG]: Dùng SUMIFS quét chính xác khối lượng từ sheet 'Hao phí VT' 
                        # dựa vào điều kiện chéo: Mã công tác (Cột B) và Tên thiết bị (Cột E)
                        hp_qty_formula = f"=SUMIFS('Hao phí VT'!$I$6:$I${last_hp_row}, 'Hao phí VT'!$B$6:$B${last_hp_row}, \"{job.code}\", 'Hao phí VT'!$E$6:$E${last_hp_row}, \"{detail_item.name}\")"
                        
                        usage_jobs.append({
                            "job": job,
                            "detail_item": detail_item,
                            "qty_formula": hp_qty_formula
                        })

        # Xác định chính xác vị trí dòng cha trước khi append
        parent_row = curr_row

        # --- A. GHI DÒNG CHA (MÁY THI CÔNG CHÍNH) ---
        row_data = [""] * 17
        row_data[0] = idx                          # A: STT (Đánh số từ 1 đến n)
        row_data[1] = m_id                         # B: Mã hiệu định danh
        row_data[4] = m_name                       # E: Tên máy thi công chính
        row_data[7] = m_unit                       # H: Đơn vị
        row_data[9] = 0                            # J: Khối lượng (Sẽ được gán hàm SUM tổng các dòng con)
        row_data[10] = 0                           # K: Giá gốc
        row_data[11] = 0                           # L: Giá thông báo
        row_data[12] = 0                           # M: Bù giá
        row_data[13] = price                       # N: Giá hiện tại
        row_data[14] = f"=N{parent_row}-K{parent_row}" # O: Chênh lệch
        row_data[15] = f"=ROUND(J{parent_row}*O{parent_row}, 0)" # P: Tổng chênh
        row_data[16] = f"=ROUND(J{parent_row}*N{parent_row}, 0)" # Q: Thành tiền

        ws_m.append(row_data)
        ws_m.cell(parent_row, 5).font = font_data_11b  # Định dạng in đậm tên thiết bị chính
        curr_row += 1

        # --- B. GHI CÁC DÒNG CON (TRUY VẾT CÔNG TÁC MÀU XANH) ---
        child_start_row = curr_row
        for item_data in usage_jobs:
            job = item_data["job"]
            detail_item = item_data["detail_item"]
            j_qty = item_data.get("qty_formula", f"=J{curr_row-1}")

            # Xác định dòng bên sheet 'Công trình'
            raw_stt = getattr(job, 'stt', 1)
            try:
                stt_int = int(raw_stt)
            except (ValueError, TypeError):
                stt_int = 1
                
            cong_trinh_row = getattr(job, 'row_index_in_cong_trinh', None) 
            if not cong_trinh_row:
                cong_trinh_row = 6 + stt_int - 1

            # 1. Cột E: Link tên công tác trực tiếp từ cột D của sheet 'Công trình'
            j_name_formula = f"='Công trình'!D{cong_trinh_row}"
            
            # 2. Cột B: Link công thức động STT công tác trực tiếp từ cột A của sheet 'Công trình'
            j_stt_formula = f"='Công trình'!A{cong_trinh_row}"

            child_row = [""] * 17
            child_row[0] = ""                      # Cột A: STT để trống cho dòng con (Chuẩn F1)
            child_row[1] = j_stt_formula           # Cột B: Link STT từ 'Công trình'!A{row}
            child_row[4] = j_name_formula          # Cột E: Trỏ động tên công tác sang sheet Công trình
            child_row[9] = j_qty                   # Cột J: Khối lượng link động qua SUMIFS
            child_row[15] = f"=ROUND(J{curr_row}*O{curr_row}, 0)"
            child_row[16] = f"=ROUND(J{curr_row}*N{curr_row}, 0)"

            ws_m.append(child_row)
            
            # Thiết lập định dạng màu xanh chuyên biệt cho các dòng con truy vết
            blue_font = Font(name="Times New Roman", size=11, bold=False, color="0000FF")
            for c_idx in range(1, 18):
                ws_m.cell(curr_row, c_idx).font = blue_font
            curr_row += 1
            
        child_end_row = curr_row - 1

        # --- C. GỘP NHÓM (GROUP) VÀ GÁN HÀM SUM CHO DÒNG CHA ---
        if child_start_row <= child_end_row:
            # Thu gọn/mở rộng dòng con theo chuẩn outline grouping của Excel
            ws_m.row_dimensions.group(child_start_row, child_end_row, outline_level=1, hidden=True)
            # [CƠ CHẾ SUM]: Gán hàm SUM tổng khối lượng các dòng con vào cột J của dòng cha
            ws_m.cell(parent_row, 10).value = f"=SUM(J{child_start_row}:J{child_end_row})"

    last_m_item_row = curr_row - 1

    # 4. GHI DÒNG CỘNG TỔNG CHI PHÍ MÁY THI CÔNG Ở CUỐI BẢNG
    total_row = [""] * 17
    total_row[2] = "Cộng Máy:"
    total_row[15] = f"=SUM(P{m_start_row}:P{last_m_item_row})"
    total_row[16] = f"=SUM(Q{m_start_row}:Q{last_m_item_row})"
    ws_m.append(total_row)
    total_m_row = curr_row

    # 5. THIẾT LẬP ĐỊNH DẠNG & KIỂU CHỮ TOÀN SHEET
    ws_m.cell(1, 1).font = font_title_15b
    ws_m.cell(2, 1).font = font_sub_11i
    ws_m.cell(3, 1).font = font_sub_11i
    ws_m.cell(4, 1).font = font_sub_11i

    # Định dạng vùng Header bảng (Dòng 6)
    for c in range(1, 18):
        cell = ws_m.cell(6, c)
        cell.font = font_header_11b
        cell.fill = fill_header_gray
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Vòng lặp áp dụng style chi tiết cho từng ô dữ liệu và tổng kết
    for r in range(m_start_row, total_m_row + 1):
        is_total = (r == total_m_row)
        for c in range(1, 18):
            cell = ws_m.cell(r, c)
            # Giữ nguyên font chữ màu xanh đối với các dòng con truy vết
            if r != total_m_row and cell.font.color and cell.font.color.rgb == "0000FF":
                pass 
            else:
                cell.font = font_data_11b if is_total else font_data_11r
            cell.border = border_total if is_total else border_cell
            if is_total:
                cell.fill = fill_sub_gray
                if c in (16, 17):
                    cell.font = Font(name="Times New Roman", size=11, bold=True, color="0000FF")
            # Định dạng kiểu số tiền (currency) và định dạng chuẩn (norm)
            if c in (11, 12, 13, 14, 15, 16, 17):
                cell.number_format = FMT_CURRENCY
            if c == 10:
                cell.number_format = FMT_NORM

    return ws_m, total_m_row