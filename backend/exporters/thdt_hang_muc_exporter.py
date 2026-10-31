# backend/exporters/thdt_hang_muc_exporter.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border
from typing import List, Dict, Any

def _init_thdt_hang_muc_sheet(wb) -> openpyxl.worksheet.worksheet.Worksheet:
    """Khởi tạo hoặc lấy sẵn sheet 'THDT hạng mục' và thiết lập hiển thị lưới."""
    sheet_name = "THDT hạng mục"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    return ws


def _write_thdt_header(ws, font_title_15b: Font, font_header_11b: Font, fill_header_gray: PatternFill, border_cell: Border):
    """Ghi tiêu đề chính và hệ thống tiêu đề cột (Header) cho sheet THDT hạng mục."""
    ws.append(["BẢNG TỔNG HỢP CHI PHÍ DỰ TOÁN HẠNG MỤC"])
    ws.append(["='Công trình'!A3"])
    ws.append(["='Công trình'!A4"])
    ws.append([]) # Dòng 4 trống

    headers = [
        "STT", "Mã hiệu chi phí", "Nội dung chi phí các khoản mục", 
        "Đơn vị", "Cách tính / Hệ số", "Thành tiền (VNĐ)", "Căn cứ pháp lý / Ghi chú"
    ]
    ws.append(headers)
    ws.row_dimensions[5].height = 25

    for c in range(1, len(headers) + 1):
        cell = ws.cell(5, c)
        cell.font = font_header_11b
        cell.fill = fill_header_gray
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell


def _write_thdt_data(
    ws, 
    items: List[Dict[str, Any]], 
    ctx: Any, 
    tot_ct_row: int, 
    last_hp_row: int, 
    font_data_11r: Font, 
    font_data_11b: Font, 
    border_cell: Border, 
    FMT_MONEY: str,
    border_total: Border = None
) -> int:
    """Ghi dữ liệu và tự động sinh dải ô công thức liên kết động từ sheet Công trình, Khoản 4.5 và Hệ số."""
    start_row = 6
    hp_end = last_hp_row if last_hp_row and last_hp_row >= 7 else 100
    
    # Lấy tọa độ dòng tổng cộng của sheet Công trình từ ctx (nếu có)
    actual_tot_ct = ctx.tot_ct_row if ctx and getattr(ctx, "tot_ct_row", None) else tot_ct_row

    # Nếu có actual_tot_ct từ sheet Công trình thì liên kết trực tiếp vào 4 cột thành tiền chi tiết R, S, T, U
    if actual_tot_ct:
        f_vl = f"='Công trình'!R{actual_tot_ct}"
        f_nc = f"='Công trình'!S{actual_tot_ct}"
        f_dc = f"='Công trình'!T{actual_tot_ct}"
        f_m  = f"='Công trình'!U{actual_tot_ct}"
        note_direct = "Tổng hợp từ bảng khối lượng công trình"
    else:
        f_vl = f"=SUM('Hao phí VT'!M7:M{hp_end})"
        f_nc = f"=SUM('Hao phí VT'!N7:N{hp_end})"
        f_dc = f"=SUM('Hao phí VT'!O7:O{hp_end})"
        f_m  = f"=SUM('Hao phí VT'!P7:P{hp_end})"
        note_direct = "Tổng hợp từ chi tiết hao phí"

    tot_45 = ctx.get_khoan_4_5_row(24) if ctx and hasattr(ctx, "get_khoan_4_5_row") else 24
    f_45 = f"='Khoản 4.5 TT20'!$J${tot_45}"

    raw_items = items if items else [
        {"ma": "CP_VL", "ten": "Chi phí vật liệu", "dvt": "Lô", "cach_tinh": "Từ sheet 'Công trình'", "formula": f_vl, "ghi_chu": note_direct},
        {"ma": "CP_NC", "ten": "Chi phí nhân công", "dvt": "Lô", "cach_tinh": "Từ sheet 'Công trình'", "formula": f_nc, "ghi_chu": note_direct},
        {"ma": "CP_DC", "ten": "Chi phí công cụ dụng cụ", "dvt": "Lô", "cach_tinh": "Từ sheet 'Công trình'", "formula": f_dc, "ghi_chu": note_direct},
        {"ma": "CP_M", "ten": "Chi phí máy và thiết bị thi công", "dvt": "Lô", "cach_tinh": "Từ sheet 'Công trình'", "formula": f_m, "ghi_chu": note_direct},
        {"ma": "CP_TC", "ten": "1. Chi phí trực tiếp (T = VL + NC + DC + M)", "dvt": "Lô", "cach_tinh": "T = VL + NC + DC + M", "is_bold": True, "ghi_chu": "Cộng chi phí trực tiếp"},
        {"ma": "CP_C", "ten": "2. Chi phí quản lý chung (C)", "dvt": "Lô", "cach_tinh": "Theo sheet 'Hệ số'", "is_bold": False, "ghi_chu": "Thông tư 02/2017/TT-BTC (Phụ lục 03/04)"},
        {"ma": "CP_KHAC", "ten": "3. Chi phí khác (Chi phí phụ trợ theo PAKT)", "dvt": "Lô", "cach_tinh": "Khoản 4.5 TT20", "formula": f_45, "ghi_chu": "Tổng hợp từ sheet Khoản 4.5 TT20 (PAKT 2026)"},
        {"ma": "CP_TN", "ten": "4. Thu nhập chịu thuế tính trước (TL)", "dvt": "Lô", "cach_tinh": "Theo sheet 'Hệ số'", "is_bold": False, "ghi_chu": "Lợi nhuận định mức (5.5%)"},
        {"ma": "CP_G", "ten": "5. Giá trị dự toán hạng mục trước thuế (G)", "dvt": "Lô", "cach_tinh": "G = T + C + Khác + TL", "is_bold": True, "is_grand_total": True, "ghi_chu": "Tổng chi phí dự toán trước thuế"}
    ]

    # Ánh xạ mã định danh để tự động tra cứu tọa độ dòng
    code_to_row = {}
    for idx, item in enumerate(raw_items):
        code_to_row[item["ma"]] = start_row + idx

    # Tra cứu dòng tỷ lệ hệ số từ ctx
    row_qlc = ctx.get_he_so_row("CP_QLC", 6) if ctx and hasattr(ctx, "get_he_so_row") else None
    row_tntt = ctx.get_he_so_row("TN_TT", 7) if ctx and hasattr(ctx, "get_he_so_row") else None

    # Sinh công thức động theo từ điển mã
    for item in raw_items:
        ma = item["ma"]
        if ma == "CP_TC":
            item["formula"] = f"=SUM(F{code_to_row['CP_VL']}:F{code_to_row['CP_M']})"
        elif ma == "CP_C":
            if row_qlc:
                item["formula"] = f"=F{code_to_row['CP_TC']}*'Hệ số'!$E${row_qlc}"
                item["cach_tinh"] = f"Tỷ lệ ('Hệ số'!E{row_qlc}) x T"
            else:
                item["formula"] = f"=F{code_to_row['CP_TC']}*0.20"
                item["cach_tinh"] = "20% x T"
        elif ma == "CP_TN":
            if row_tntt:
                item["formula"] = f"=(F{code_to_row['CP_TC']}+F{code_to_row['CP_C']}+F{code_to_row['CP_KHAC']})*'Hệ số'!$E${row_tntt}"
                item["cach_tinh"] = f"Tỷ lệ ('Hệ số'!E{row_tntt}) x (T + C + Khác)"
            else:
                item["formula"] = f"=(F{code_to_row['CP_TC']}+F{code_to_row['CP_C']}+F{code_to_row['CP_KHAC']})*0.055"
                item["cach_tinh"] = "5.5% x (T + C + Khác)"
        elif ma == "CP_G":
            item["formula"] = f"=F{code_to_row['CP_TC']}+F{code_to_row['CP_C']}+F{code_to_row['CP_KHAC']}+F{code_to_row['CP_TN']}"

    curr_row = start_row
    for idx, item in enumerate(raw_items, 1):
        is_bold = item.get("is_bold", False)
        is_grand = item.get("is_grand_total", False)
        cach_tinh_val = str(item.get("cach_tinh", ""))
        # Đảm bảo cột mô tả cách tính luôn là chuỗi thuần túy, không để openpyxl nhận diện thành công thức Excel
        if cach_tinh_val.startswith("="):
            cach_tinh_val = cach_tinh_val.lstrip("=")

        ws.append([
            idx if not is_bold else "",
            item.get("ma", ""),
            item.get("ten", ""),
            item.get("dvt", "Lô"),
            cach_tinh_val,
            item.get("formula", 0),
            item.get("ghi_chu", "")
        ])

        for c in range(1, 8):
            cell = ws.cell(curr_row, c)
            cell.font = font_data_11b if is_bold else font_data_11r
            cell.border = border_total if (is_grand and border_total) else border_cell
            if c in (1, 2, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (3, 5, 7):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = FMT_MONEY

        curr_row += 1

    return curr_row - 1


def _apply_thdt_column_widths(ws):
    col_widths = {
        'A': 8, 'B': 18, 'C': 52, 'D': 10, 'E': 28, 'F': 24, 'G': 38
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width


def build_thdt_hang_muc_sheet(
    wb, 
    ctx: Any = None,
    items: List[Dict[str, Any]] = None,
    tot_ct_row: int = None,
    last_hp_row: int = 100,
    font_title_15b: Font = None, 
    font_header_11b: Font = None, 
    font_data_11r: Font = None, 
    font_data_11b: Font = None,
    fill_header_gray: PatternFill = None, 
    border_cell: Border = None, 
    border_total: Border = None,
    FMT_MONEY: str = "#,##0"
):
    """
    Hàm điều phối chính: Nhận đối tượng ctx (ProjectContext) để trỏ công thức động 100%.
    """
    ws = _init_thdt_hang_muc_sheet(wb)
    _write_thdt_header(ws, font_title_15b, font_header_11b, fill_header_gray, border_cell)
    last_row = _write_thdt_data(ws, items, ctx, tot_ct_row, last_hp_row, font_data_11r, font_data_11b, border_cell, FMT_MONEY, border_total=border_total)
    _apply_thdt_column_widths(ws)

    return ws, last_row