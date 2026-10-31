import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Tuple
from openpyxl.worksheet.worksheet import Worksheet


def _build_cong_trinh_header(
    ws_ct: Worksheet, 
    font_title_15b: Font, 
    font_sub_11i: Font, 
    font_data_11b: Font, 
    font_header_11b: Font, 
    fill_header_gray: PatternFill
) -> None:
    """1. Tạo phần tiêu đề công trình và hàng header 24 cột chuẩn F1."""
    ws_ct.append(["BẢNG CHI TIẾT KHỐI LƯỢNG CÔNG TRÌNH QUAN TRẮC MÔI TRƯỜNG"])
    ws_ct.append(["(Căn cứ theo Quyết định số 1155/QĐ-NĐVT4 và Thông tư số 20/2017/TT-BTNMT)"])
    ws_ct.append(["CÔNG TRÌNH: QUAN TRẮC MÔI TRƯỜNG CHO NMNĐ VĨNH TÂN 4 VÀ VĨNH TÂN 4 MỞ RỘNG NĂM 2026"])
    ws_ct.append(["HẠNG MỤC: TOÀN BỘ KHỐI LƯỢNG QUAN TRẮC MÔI TRƯỜNG THEO ĐỊNH MỨC THÔNG TƯ 20/2017/TT-BTNMT"])
    ws_ct.append([])
    ws_ct.append([
        "STT", "Ký hiệu vị trí (Bản vẽ)", "Mã hiệu công tác (TT20)", "Danh mục công tác quan trắc", "Đơn vị",
        "Số vị trí quan trắc", "Số đợt / Tần suất (lần/năm)", "Số mẫu song song", "Hệ số chung mẫu (M)",
        "Kích thước phụ", "Ghi chú phương pháp", "KHỐI LƯỢNG THỰC HIỆN (Q)",
        "Đơn giá Vật liệu (Đ/mẫu)", "Đơn giá Nhân công (Đ/mẫu)", "Đơn giá Dụng cụ (Đ/mẫu)", "Đơn giá Máy thi công (Đ/mẫu)",
        "CỘNG ĐƠN GIÁ TRỰC TIẾP (Đ/mẫu)",
        "THÀNH TIỀN VẬT LIỆU (VNĐ)", "THÀNH TIỀN NHÂN CÔNG (VNĐ)", "THÀNH TIỀN DỤNG CỤ (VNĐ)", "THÀNH TIỀN MÁY THI CÔNG (VNĐ)",
        "CỘNG THÀNH TIỀN TRỰC TIẾP (VNĐ)",
        "ĐƠN GIÁ DỰ THẦU TRỌN GÓI (Đ/mẫu)", "THÀNH TIỀN DỰ THẦU TRỌN GÓI (VNĐ)"
    ])

    ws_ct.cell(1, 1).font = font_title_15b
    ws_ct.cell(2, 1).font = font_sub_11i
    ws_ct.cell(3, 1).font = font_data_11b
    ws_ct.cell(4, 1).font = font_data_11b

    for c in range(1, 25):
        cell = ws_ct.cell(6, c)
        cell.font = font_header_11b
        cell.fill = fill_header_gray
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_cong_trinh_section(
    ws_ct: Worksheet, 
    r: int, 
    curr_section: str, 
    font_sec_12b: Font, 
    fill_header_gray: PatternFill
) -> None:
    """2. Ghi dòng tiêu đề phân đoạn (Phần: Giám sát môi trường...)."""
    ws_ct.append(["", "", "", f"PHẦN: GIÁM SÁT {curr_section.upper()}"] + [""] * 20)
    sec_cell = ws_ct.cell(r, 4)
    sec_cell.font = font_sec_12b
    for c in range(1, 25):
        ws_ct.cell(r, c).fill = fill_header_gray


def _write_cong_trinh_row(
    ws_ct: Worksheet, 
    idx: int, 
    r: int, 
    it: Dict[str, Any], 
    map_dgct: Dict[Any, Dict[str, Any]]
) -> None:
    so_vt = it.get("so_vi_tri", 1) or 1
    so_lan = it.get("so_lan", 1) or 1
    vi_tri_str = it.get("vi_tri", f"VT_{idx}")
    code_str = it.get("ma_gop", f"TT20_{idx:03d}")
    name_str = it.get("thong_so", "Thông số quan trắc")
    dvt_str = it.get("dvt", "mẫu")
    m_val = it.get("so_thong_so_chung", 1) or 1

    f_q = f"=F{r}*G{r}"

    dgct_info = map_dgct.get(idx) or map_dgct.get(str(idx), {})
    row_vl = dgct_info.get('row_vl')
    row_nc = dgct_info.get('row_nc')
    row_dc = dgct_info.get('row_dc')
    row_m  = dgct_info.get('row_m')
    row_total = dgct_info.get('row_total')

    # Đơn giá từng cấu phần
    f_vl = f"='Đơn giá chi tiết'!$J${row_vl}" if row_vl else "0"
    f_nc = f"='Đơn giá chi tiết'!$J${row_nc}" if row_nc else "0"
    f_dc = f"='Đơn giá chi tiết'!$J${row_dc}" if row_dc else "0"
    f_m  = f"='Đơn giá chi tiết'!$J${row_m}"  if row_m  else "0"

    f_unit = f"=M{r}+N{r}+O{r}+P{r}"

    # Thành tiền từng cấu phần trực tiếp theo đúng chỉ đạo
    f_tt_vl = f"=ROUND(L{r}*M{r}, 0)"
    f_tt_nc = f"=ROUND(L{r}*N{r}, 0)"
    f_tt_dc = f"=ROUND(L{r}*O{r}, 0)"
    f_tt_m  = f"=ROUND(L{r}*P{r}, 0)"

    f_total = f"=ROUND(L{r}*Q{r}, 0)"

    f_all_in = f"='Đơn giá chi tiết'!$J${row_total}" if row_total else "0"
    f_total_all_in = f"=ROUND(L{r}*W{r}, 0)"

    ws_ct.append([
        idx, vi_tri_str, code_str, name_str, dvt_str,
        so_vt, so_lan, 1, m_val,
        "", it.get("sec_id", "TT20"), f_q,
        f_vl, f_nc, f_dc, f_m,
        f_unit,
        f_tt_vl, f_tt_nc, f_tt_dc, f_tt_m,
        f_total,
        f_all_in, f_total_all_in
    ])


def _write_cong_trinh_total(
    ws_ct: Worksheet, 
    tot_ct_row: int, 
    start_ct_row: int, 
    last_ct_row: int, 
    font_data_11b: Font, 
    border_total: Border, 
    fill_total_pale: PatternFill, 
    FMT_CURRENCY: str
) -> None:
    """4. Ghi dòng tổng cộng toàn bộ công trình quan trắc với đầy đủ tổng các cấu phần."""
    ws_ct.append([
        "", "", "", "CỘNG TOÀN BỘ CÔNG TRÌNH QUAN TRẮC MÔI TRƯỜNG", "",
        "", "", "", "", "", "", "", "", "", "", "", "",
        f"=SUM(R{start_ct_row}:R{last_ct_row})", # R: Tổng Thành tiền Vật liệu
        f"=SUM(S{start_ct_row}:S{last_ct_row})", # S: Tổng Thành tiền Nhân công
        f"=SUM(T{start_ct_row}:T{last_ct_row})", # T: Tổng Thành tiền CCDC
        f"=SUM(U{start_ct_row}:U{last_ct_row})", # U: Tổng Thành tiền Máy thi công
        f"=SUM(V{start_ct_row}:V{last_ct_row})", # V: Tổng Thành tiền Trực tiếp
        "",                                      # W: Đơn giá trọn gói
        f"=SUM(X{start_ct_row}:X{last_ct_row})"  # X: Tổng Thành tiền Trọn gói
    ])

    for c in range(1, 25):
        cell = ws_ct.cell(tot_ct_row, c)
        cell.font = font_data_11b
        cell.border = border_total
        if c in (18, 19, 20, 21, 22, 24):
            cell.fill = fill_total_pale
            cell.number_format = FMT_CURRENCY


def _apply_cong_trinh_styling(
    ws_ct: Worksheet, 
    start_ct_row: int, 
    last_ct_row: int, 
    font_sec_12b: Font, 
    font_data_11r: Font, 
    border_cell: Border, 
    FMT_CURRENCY: str
) -> None:
    """5. Áp dụng định dạng style cho các dòng dữ liệu công trình."""
    for r in range(start_ct_row, last_ct_row + 1):
        for c in range(1, 25):
            cell = ws_ct.cell(r, c)
            if cell.font != font_sec_12b:
                cell.font = font_data_11r
                cell.border = border_cell
            if c in (1, 5, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c in (12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24):
                cell.number_format = FMT_CURRENCY


def _apply_cong_trinh_column_widths(ws_ct: Worksheet) -> None:
    """Thiết lập độ rộng cột thoáng chuẩn cho Sheet Công trình."""
    col_widths = {
        'A': 8, 'B': 18, 'C': 22, 'D': 42, 'E': 10,
        'F': 12, 'G': 12, 'H': 10, 'I': 12, 'J': 12,
        'K': 16, 'L': 16, 'M': 16, 'N': 16, 'O': 16,
        'P': 16, 'Q': 18, 'R': 18, 'S': 18, 'T': 18,
        'U': 18, 'V': 20, 'W': 18, 'X': 20
    }
    for col_letter, width in col_widths.items():
        ws_ct.column_dimensions[col_letter].width = width


def update_cong_trinh_from_dgct(
    ws_ct: Worksheet, 
    item_ct_rows: Dict[int, int], 
    map_dgct: Dict[Any, Dict[str, Any]]
) -> None:
    """
    6. Cập nhật các cột M -> X của sheet 'Công trình' 
    sau khi đã sinh xong sheet 'Đơn giá chi tiết' và có tọa độ chính xác.
    """
    for idx, r in item_ct_rows.items():
        dgct_info = map_dgct.get(idx) or map_dgct.get(str(idx), {})
        row_vl = dgct_info.get('row_vl')
        row_nc = dgct_info.get('row_nc')
        row_dc = dgct_info.get('row_dc')
        row_m  = dgct_info.get('row_m')
        row_total = dgct_info.get('row_total')

        # Đơn giá 4 thành phần & Đơn giá trực tiếp
        ws_ct.cell(r, 13).value = f"='Đơn giá chi tiết'!$J${row_vl}" if row_vl else "0"
        ws_ct.cell(r, 14).value = f"='Đơn giá chi tiết'!$J${row_nc}" if row_nc else "0"
        ws_ct.cell(r, 15).value = f"='Đơn giá chi tiết'!$J${row_dc}" if row_dc else "0"
        ws_ct.cell(r, 16).value = f"='Đơn giá chi tiết'!$J${row_m}" if row_m else "0"
        ws_ct.cell(r, 17).value = f"=M{r}+N{r}+O{r}+P{r}"

        # Thành tiền 4 cấu phần trực tiếp theo chỉ đạo
        ws_ct.cell(r, 18).value = f"=ROUND(L{r}*M{r}, 0)"
        ws_ct.cell(r, 19).value = f"=ROUND(L{r}*N{r}, 0)"
        ws_ct.cell(r, 20).value = f"=ROUND(L{r}*O{r}, 0)"
        ws_ct.cell(r, 21).value = f"=ROUND(L{r}*P{r}, 0)"
        ws_ct.cell(r, 22).value = f"=ROUND(L{r}*Q{r}, 0)"

        # Đơn giá và thành tiền trọn gói
        ws_ct.cell(r, 23).value = f"='Đơn giá chi tiết'!$J${row_total}" if row_total else "0"
        ws_ct.cell(r, 24).value = f"=ROUND(L{r}*W{r}, 0)"


def build_cong_trinh_sheet(
    wb, 
    pakt_items: List[Dict[str, Any]], 
    map_dgct: Dict[Any, Dict[str, Any]], 
    font_title_15b: Font, 
    font_sub_11i: Font, 
    font_header_11b: Font, 
    font_sec_12b: Font, 
    font_data_11r: Font, 
    font_data_11b: Font, 
    fill_header_gray: PatternFill, 
    fill_sub_gray: PatternFill, 
    fill_total_pale: PatternFill, 
    border_cell: Border, 
    border_total: Border, 
    FMT_CURRENCY: str
) -> Tuple[Worksheet, Dict[int, int], int]:
    """
    7. Xây dựng Sheet 'Công trình' và trả về tuple: (worksheet, item_ct_rows, tot_ct_row).
    item_ct_rows là registry lưu tọa độ dòng Excel thực tế của từng công tác (idx -> row_excel).
    tot_ct_row là tọa độ dòng tổng cộng của sheet Công trình.
    """
    if "Công trình" in wb.sheetnames:
        ws_ct = wb["Công trình"]
        wb.remove(ws_ct)
        
    ws_ct = wb.create_sheet(title="Công trình")
    ws_ct.views.sheetView[0].showGridLines = True
    
    _build_cong_trinh_header(ws_ct, font_title_15b, font_sub_11i, font_data_11b, font_header_11b, fill_header_gray)

    start_ct_row = 7
    curr_section = ""
    row_offset = 0
    item_ct_rows: Dict[int, int] = {}

    for idx, it in enumerate(pakt_items, start=1):
        r = start_ct_row + row_offset
        moi_truong = it.get("moi_truong", "KHÔNG KHÍ XUNG QUANH")

        if moi_truong != curr_section:
            curr_section = moi_truong
            _write_cong_trinh_section(ws_ct, r, curr_section, font_sec_12b, fill_header_gray)
            row_offset += 1
            r = start_ct_row + row_offset

        item_ct_rows[idx] = r
        _write_cong_trinh_row(ws_ct, idx, r, it, map_dgct)
        row_offset += 1

    last_ct_row = start_ct_row + row_offset - 1
    tot_ct_row = last_ct_row + 1
    
    _write_cong_trinh_total(ws_ct, tot_ct_row, start_ct_row, last_ct_row, font_data_11b, border_total, fill_total_pale, FMT_CURRENCY)
    _apply_cong_trinh_styling(ws_ct, start_ct_row, last_ct_row, font_sec_12b, font_data_11r, border_cell, FMT_CURRENCY)
    _apply_cong_trinh_column_widths(ws_ct)
    ws_ct.freeze_panes = 'A7'

    return ws_ct, item_ct_rows, tot_ct_row
