# backend/exporters/khoan_4_5_tt20_exporter.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border
from typing import List, Dict, Any

def _init_khoan_4_5_sheet(wb) -> openpyxl.worksheet.worksheet.Worksheet:
    """Khởi tạo hoặc lấy sẵn sheet 'Khoản 4.5 TT20' và thiết lập hiển thị lưới."""
    sheet_name = "Khoản 4.5 TT20"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    return ws


def _write_khoan_4_5_header(ws, font_title_15b: Font, font_header_11b: Font, fill_header_gray: PatternFill, border_cell: Border):
    """Ghi tiêu đề chính và hệ thống tiêu đề cột (Header) cho sheet Khoản 4.5 TT20."""
    ws.append(["TÍNH TOÁN CHI PHÍ PHỤ TRỢ VÀ KHOẢN 4.5 THEO THÔNG TƯ 20/2017/TT-BTNMT"])
    ws.append(["='Công trình'!A3"])
    ws.append(["='Công trình'!A4"])
    ws.append([]) # Dòng 4 trống

    # Tiêu đề cột chuẩn bảng chi phí
    headers = [
        "STT", "Mã hiệu / Nội dung chi phí", "Diễn giải chi tiết nội dung tính toán theo PAKT", 
        "Đơn vị", "Giá trị / Chi phí gốc (VNĐ)", "Tỷ lệ (%) / Hệ số", "Thành tiền (VNĐ)", "Căn cứ pháp lý / Ghi chú"
    ]
    ws.append(headers)
    ws.row_dimensions[5].height = 28

    # Định dạng dòng header (Dòng 5)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(5, c)
        cell.font = font_header_11b
        cell.fill = fill_header_gray
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell


def _write_khoan_4_5_data(ws, items: List[Dict[str, Any]], font_data_11r: Font, border_cell: Border, FMT_MONEY: str) -> int:
    """Ghi dữ liệu các khoản chi phí phụ trợ bám sát PAKT vào sheet."""
    start_row = 6
    curr_row = start_row

    # Dữ liệu mặc định bám sát Phương án kỹ thuật (PAKT)
    default_items = items if items else [
        
        {
            "ma": "CP_DI_CHUYEN", 
            "ten": "Chi phí di chuyển tuyến xa (Hà Nội - Vĩnh Tân)", 
            "diễn_giải": "Cự ly khứ hồi 2.800 km/chuyến bằng ô tô cho tổng số công di chuyển", 
            "dvt": "Chuyến", 
            "base": 14000000, 
            "rate": 4.0, 
            "ghi_chu": "Bám sát Phụ lục 2 - PAKT 2026"
        },
        {
            "ma": "CP_LUU_TRU", 
            "ten": "Chi phí lưu trú nhân sự hiện trường", 
            "diễn_giải": "Số đêm lưu trú thực tế các đợt (Đợt 2, 4: 25 đêm; Đợt 1, 3: 6 đêm đối chứng)", 
            "dvt": "Đêm", 
            "base": 620000, 
            "rate": 62.0, 
            "ghi_chu": "Bám sát Mục 2.1.2 PAKT"
        },
        {
            "ma": "CP_VAN_CHUYEN_MAU", 
            "ten": "Chi phí tách chuyến vận chuyển mẫu gấp", 
            "diễn_giải": "Vận chuyển mẫu có thời gian bảo quản ngắn (Coliform, Crom VI, BOD5, Sunfua)", 
            "dvt": "Chuyến", 
            "base": 2500000, 
            "rate": 4.0, 
            "ghi_chu": "Yêu cầu kỹ thuật bảo quản mẫu"
        },
        {
            "ma": "CP_IN_AN", 
            "ten": "Chi phí văn phòng phẩm, in ấn và lập báo cáo", 
            "diễn_giải": "In ấn kết quả phân tích và đóng quyển hồ sơ báo cáo cho 4 đợt quan trắc", 
            "dvt": "Đợt", 
            "base": 500000, 
            "rate": 4.0, 
            "ghi_chu": "Khoán gọn 500.000đ/đợt x 4 đợt"
        }
    ]

    for idx, item in enumerate(default_items, 1):
        ma = item.get("ma", "")
        rate = item.get("rate", 0)
        
        ws.append([
            idx,
            ma,
            item.get("diễn_giải", item.get("ten", "")),
            item.get("dvt", "Lô"),
            item.get("base", 0),
            rate,
            f"=E{curr_row}*F{curr_row}", # Công thức thành tiền = Gốc * Tỷ lệ/Hệ số
            item.get("ghi_chu", "Phương án kỹ thuật 2026")
        ])

        # Định dạng ô dữ liệu
        for c in range(1, 9):
            cell = ws.cell(curr_row, c)
            cell.font = font_data_11r
            cell.border = border_cell
            if c in (1, 2, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c in (5, 7):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = FMT_MONEY
            elif c == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                # Nếu là hệ số số lượng chuyến/đêm/đợt thì hiển thị dạng thập phân hoặc số nguyên, nếu là phần trăm thì hiển thị %
                if ma in ["CP_DI_CHUYEN", "CP_LUU_TRU", "CP_VAN_CHUYEN_MAU", "CP_IN_AN"]:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "0.00%"

        curr_row += 1

    return curr_row - 1


def _apply_khoan_4_5_column_widths(ws):
    """Thiết lập độ rộng chuẩn thoáng cho các cột."""
    col_widths = {
        'A': 8, 'B': 24, 'C': 48, 'D': 12, 'E': 20, 'F': 14, 'G': 22, 'H': 32
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width


def build_khoan_4_5_tt20_sheet(
    wb, 
    items: List[Dict[str, Any]] = None,
    font_title_15b: Font = None, 
    font_header_11b: Font = None, 
    font_data_11r: Font = None, 
    fill_header_gray: PatternFill = None, 
    border_cell: Border = None, 
    FMT_MONEY: str = "#,##0"
):
    """
    Hàm điều phối chính: Xây dựng hoàn chỉnh sheet 'Khoản 4.5 TT20' bám sát PAKT.
    """
    ws = _init_khoan_4_5_sheet(wb)
    _write_khoan_4_5_header(ws, font_title_15b, font_header_11b, fill_header_gray, border_cell)
    last_row = _write_khoan_4_5_data(ws, items, font_data_11r, border_cell, FMT_MONEY)
    _apply_khoan_4_5_column_widths(ws)

    return ws, last_row