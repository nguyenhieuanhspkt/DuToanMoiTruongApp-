import openpyxl
import sys
import os
import json
import re

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"D:\onedrive_hieuna\OneDrive - EVN\Tổ Thẩm định\Năm 2026\Thẩm định 138_hieuna_3\2026-04-03-HSDT\vn0900233631_Bảng giá (1).xlsx"
print(f"Loading Excel: {excel_path}")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb["Mẫu số 12.1B. Bảng giá dự thầu "]

items_kqlcnt = []
for r in range(5, ws.max_row + 1):
    stt = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    code = ws.cell(r, 3).value
    brand = ws.cell(r, 4).value
    origin = ws.cell(r, 6).value
    tech = ws.cell(r, 8).value
    dvt = ws.cell(r, 9).value
    qty = ws.cell(r, 10).value
    price_vat = ws.cell(r, 12).value
    total_vat = ws.cell(r, 13).value
    
    if name and price_vat is not None:
        try:
            p_vat = float(price_vat)
            p_pre = round(p_vat / 1.08, 2)
            items_kqlcnt.append({
                "stt": stt,
                "name": str(name).strip(),
                "code": str(code or "").strip(),
                "brand": str(brand or "").strip(),
                "origin": str(origin or "").strip(),
                "tech": str(tech or "").strip(),
                "dvt": str(dvt or "").strip(),
                "qty": qty,
                "price_vat_8": p_vat,
                "price_pre_vat": p_pre,
                "total_vat": float(total_vat) if total_vat else 0
            })
        except Exception as e:
            pass

print(f"Total items in KQLCNT IB2600096765: {len(items_kqlcnt)}")

# Save to json for reference
out_json = r"D:\TaskApp_kiet\DuToanMoiTruongApp\backend\data\kqlcnt_ib2600096765.json"
with open(out_json, "w", encoding="utf-8") as f:
    json.dump({
        "tbmt": "IB2600096765",
        "goi_thau": "Mua sắm hóa chất, vật tư Nhà máy Nhiệt điện Vĩnh Tân 4",
        "thue_suat_vat": 0.08,
        "cong_thuc_quy_doi": "Đơn giá trước thuế = Đơn giá dự thầu (VAT 8%) / 1.08",
        "total_items": len(items_kqlcnt),
        "items": items_kqlcnt
    }, f, ensure_ascii=False, indent=2)

print(f"Saved extracted items to {out_json}")

# Now load project materials and tools/equipment
mat_json_path = r"D:\TaskApp_kiet\DuToanMoiTruongApp\backend\data\material_pricing.json"
with open(mat_json_path, "r", encoding="utf-8") as f:
    mat_data = json.load(f)

proj_materials = mat_data.get("materials", [])
print(f"\nProject materials count: {len(proj_materials)}")

# Load msc_lookup_materials
msc_mat_path = r"D:\TaskApp_kiet\DuToanMoiTruongApp\backend\data\msc_lookup_materials.json"
with open(msc_mat_path, "r", encoding="utf-8") as f:
    msc_mat_data = json.load(f)

msc_materials = msc_mat_data.get("items", [])
print(f"Project MSC materials count: {len(msc_materials)}")

# Let's inspect potential matches between IB2600096765 and our materials
def normalize(s):
    if not s: return ""
    s = s.lower()
    s = re.sub(r'[\(\)\,\.\-\+\%\/]', ' ', s)
    return " ".join(s.split())

matches = []
for p in proj_materials:
    p_name = p.get("name", "")
    p_norm = normalize(p_name)
    best_m = None
    for k in items_kqlcnt:
        k_name = k.get("name", "")
        k_norm = normalize(k_name)
        # Check direct inclusion or token overlap
        p_tokens = set(p_norm.split())
        k_tokens = set(k_norm.split())
        overlap = len(p_tokens.intersection(k_tokens))
        
        # Substring checks
        if p_norm in k_norm or k_norm in p_norm or overlap >= 2:
            matches.append({
                "project_code": p.get("code"),
                "project_name": p_name,
                "project_dvt": p.get("unit"),
                "current_unit_price": p.get("unit_price"),
                "ib_stt": k["stt"],
                "ib_name": k["name"],
                "ib_dvt": k["dvt"],
                "ib_brand": k["brand"],
                "ib_origin": k["origin"],
                "ib_price_vat": k["price_vat_8"],
                "ib_price_pre_vat": k["price_pre_vat"],
                "overlap": overlap
            })

print(f"\nPotential matches found: {len(matches)}")
print("\nTop matched candidates:")
seen = set()
for m in matches:
    key = (m["project_name"], m["ib_name"])
    if key not in seen:
        seen.add(key)
        print(f"- [Dự toán]: {m['project_name']} ({m['project_dvt']}) - Đơn giá hiện tại: {m['current_unit_price']:,.0f}")
        print(f"  [IB2600096765]: {m['ib_name']} ({m['ib_dvt']}) - Brand: {m['ib_brand']} ({m['ib_origin']})")
        print(f"  -> Giá trúng thầu (VAT 8%): {m['ib_price_vat']:,.0f} | Giá trước thuế (/1.08): {m['ib_price_pre_vat']:,.2f}\n")
