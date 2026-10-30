# backend/exporters/hao_phi_vt_exporter.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any

def build_hao_phi_vt_sheet(
    wb, 
    pakt_items: List[Dict[str, Any]], 
    start_ct_row: int,
    job_mapper: Any, # <--- Nhận job_mapper ở đây
    font_title_15b: Font, 
    font_sub_11i: Font, 
    font_header_11b: Font, 
    font_data_11r: Font, 
    font_data_11b: Font, 
    fill_header_gray: PatternFill, 
    border_cell: Border, 
    FMT_NORM: str
):
    registry = job_mapper.registry if job_mapper else None
    """
    Xây dựng Sheet 'Hao phí VT' với chuẩn giao diện F1 và liên kết công thức động 
    tới tọa độ cell của sheet 'Dữ liệu định mức'.
    """
    if "Hao phí VT" in wb.sheetnames:
        ws_hp = wb["Hao phí VT"]
        ws_hp.views.sheetView[0].showGridLines = True
    else:
        ws_hp = wb.create_sheet(title="Hao phí VT")
        ws_hp.views.sheetView[0].showGridLines = True

    # Tiêu đề bảng
    ws_hp.append(["HAO PHÍ VẬT LIỆU, NHÂN CÔNG, MÁY VÀ CÔNG CỤ DỤNG CỤ CHI TIẾT THEO ĐỊNH MỨC THÔNG TƯ 20/2017/TT-BTNMT"])
    ws_hp.append(["='Công trình'!A3"])
    ws_hp.append(["='Công trình'!A4"])
    ws_hp.append([]) # Dòng 4 trống

    # -------------------------------------------------------------
    # THIẾT KẾ MULTI-LEVEL HEADER (Dòng 5 và Dòng 6)
    # -------------------------------------------------------------
    ws_hp.append([
        "STT", "Mã hiệu TT20", "Phân loại", "Tên công tác / Thành phần hao phí chi tiết", "Đơn vị", "Khối lượng công trình (Q)",
        "MỨC HAO PHÍ", "", "", "", "Hệ số", "KHỐI LƯỢNG HAO PHÍ", "", "", "", "Căn cứ pháp lý định mức TT20"
    ])
    ws_hp.append([
        "", "", "", "", "", "",
        "Vật liệu", "Nhân công", "Công cụ dụng cụ", "Máy thi công",
        "",
        "Vật liệu", "Nhân công", "Công cụ dụng cụ", "Máy thi công",
        ""
    ])

    # Gộp ô các tiêu đề cấp lớn ở Dòng 5 & Dòng 6
    ws_hp.merge_cells("A5:A6")
    ws_hp.merge_cells("B5:B6")
    ws_hp.merge_cells("C5:C6")
    ws_hp.merge_cells("D5:D6")
    ws_hp.merge_cells("E5:E6")
    ws_hp.merge_cells("F5:F6")
    
    ws_hp.merge_cells("G5:J5")  # Cụm "MỨC HAO PHÍ" (G đến J)
    ws_hp.merge_cells("K5:K6")  # Hệ số
    ws_hp.merge_cells("L5:O5")  # Cụm "KHỐI LƯỢNG HAO PHÍ" (L đến O)
    ws_hp.merge_cells("P5:P6")  # Căn cứ pháp lý

    hp_row = 7 # Dữ liệu bắt đầu từ dòng 7

    # Duyệt qua các công việc thực tế để bung chi tiết định mức
    for i, it in enumerate(pakt_items, 1):
        ct_r = start_ct_row + i - 1
        tt = it.get('tt20_data') or {}
        core_code = tt.get('coreCode') or it.get('code') or it.get('ma_gop') or 'TT20'
        work_name = it.get('thong_so') or 'Công tác quan trắc'
        work_unit = it.get('dvt') or 'mẫu'

        # Dòng Cha (Mã công tác chính)
        ws_hp.append([
            i, core_code, "CÔNG TÁC", work_name, work_unit,
            f"='Công trình'!L{ct_r}", "", "", "", "", "", "", "", "", "", "Thông tư 20/2017/TT-BTNMT"
        ])
        master_r = hp_row
        ws_hp.cell(master_r, 4).font = font_data_11b
        for c in range(1, 17):
            ws_hp.cell(master_r, c).fill = fill_header_gray
            ws_hp.cell(master_r, c).border = border_cell
        hp_row += 1

        # Bung từng giai đoạn (Ngoại nghiệp, Nội nghiệp, Toàn phần)
        for stage_key, stage_lbl in [('field', 'Ngoại nghiệp'), ('lab', 'Nội nghiệp'), ('single', 'Toàn phần')]:
            sub = tt.get(stage_key)
            if not sub:
                continue

            # 1. NHÂN CÔNG (Cột H: Mức NC, Cột M: Khối lượng NC)
            nc_norm = sub.get('dinh_muc_cong')
            nc_title = sub.get('dinh_bien') or f'Quan trắc viên ({stage_lbl})'
            if nc_norm is not None and float(nc_norm) > 0:
                uid_nc = registry.generate_unique_id(core_code, 'NC', 1) if registry else None
                norm_val_nc = registry.get_formula_by_id(uid_nc, float(nc_norm)) if registry else float(nc_norm)
                
                ws_hp.append([
                    "", ("NC_LDPT" if ("phổ thông" in nc_title.lower() or "phụ việc" in nc_title.lower()) else "NC_LDKT"), 
                    "NHÂN CÔNG", f" - {nc_title} ({stage_lbl})", "công", "",
                    "", norm_val_nc, "", "", "",   # Mức hao phí (Cột H = Nhân công)
                    "", f"=ROUND(F{master_r}*H{hp_row}, 4)", "", "",  # Khối lượng hao phí (Cột M = Nhân công)
                    "Bảng 1 TT20 & TTLT 56/2015"
                ])
                hp_row += 1

            # 2. VẬT LIỆU & HÓA CHẤT (Cột G: Mức VL, Cột L: Khối lượng VL)
            for m_idx, m in enumerate(sub.get('materials', []), 1):
                m_name = m.get('name', '').strip()
                m_unit = m.get('unit', '').strip()
                m_norm = float(m.get('norm') or 0.0)
                if m_norm > 0:
                    uid_vl = registry.generate_unique_id(core_code, 'VL', m_idx) if registry else None
                    norm_val_vl = registry.get_formula_by_id(uid_vl, m_norm) if registry else m_norm

                    ws_hp.append([
                        "", f"VL_{core_code}", "VẬT LIỆU", f" - {m_name} ({stage_lbl})", m_unit, "",
                        norm_val_vl, "", "", "", "",   # Mức hao phí (Cột G = Vật liệu)
                        f"=ROUND(F{master_r}*G{hp_row}, 4)", "", "", "",  # Khối lượng VL (Cột L)
                        "Bảng 5 TT20/2017"
                    ])
                    hp_row += 1

            # 3. CÔNG CỤ DỤNG CỤ (Cột I: Mức DC, Cột N: Khối lượng DC)
            for t_idx, t in enumerate(sub.get('tools', []), 1):
                raw_name = str(t.get('name') or '').strip()
                raw_unit = str(t.get('unit') or '').strip()
                raw_life = str(t.get('lifespan_months') or '').strip().lower()
                t_norm = float(t.get('norm') or 0.0)

                if raw_name.isdigit() or raw_name == "":
                    t_name = raw_unit if raw_unit and not raw_unit.isdigit() else "Công cụ dụng cụ"
                    t_unit = raw_life if raw_life in ['cái', 'bộ', 'đôi', 'hộp', 'chiếc', 'dây', 'tấm', 'cuốn', 'bình'] else 'cái'
                else:
                    t_name = raw_name
                    t_unit = raw_unit or 'cái'

                if t_norm > 0 and t_name:
                    uid_dc = registry.generate_unique_id(core_code, 'DC', t_idx) if registry else None
                    norm_val_dc = registry.get_formula_by_id(uid_dc, t_norm) if registry else t_norm

                    ws_hp.append([
                        "", f"DC_{core_code}", "DỤNG CỤ", f" - {t_name} ({stage_lbl})", t_unit, "",
                        "", "", norm_val_dc, "", "",   # Mức hao phí (Cột I = CCDC)
                        "", "", f"=ROUND(F{master_r}*I{hp_row}, 4)", "",  # Khối lượng CCDC (Cột N)
                        "Bảng 3 TT20/2017"
                    ])
                    hp_row += 1

            # 4. MÁY THI CÔNG & THIẾT BỊ (Cột J: Mức M, Cột O: Khối lượng M)
            for e_idx, e in enumerate(sub.get('equipment', []), 1):
                e_name = e.get('name', '').strip()
                e_unit = e.get('unit', '').strip()
                e_norm = float(e.get('norm') or 0.0)
                if e_norm > 0:
                    uid_m = registry.generate_unique_id(core_code, 'M', e_idx) if registry else None
                    norm_val_m = registry.get_formula_by_id(uid_m, e_norm) if registry else e_norm

                    ws_hp.append([
                        "", f"M_{core_code}", "THIẾT BỊ", f" - {e_name} ({stage_lbl})", e_unit, "",
                        "", "", "", norm_val_m, "",   # Mức hao phí (Cột J = Máy)
                        "", "", "", f"=ROUND(F{master_r}*J{hp_row}, 4)",  # Khối lượng Máy (Cột O)
                        "Bảng 4 TT20 & TT 45/2013"
                    ])
                    hp_row += 1

        # Dòng trống ngăn cách các công tác
        ws_hp.append([])
        hp_row += 1

    last_hp_row = hp_row - 1

    # Format toàn bộ các ô dữ liệu trong sheet
    for r in range(5, last_hp_row + 1):
        for c in range(1, 17):
            cell = ws_hp.cell(r, c)
            if r in (5, 6):
                cell.font = font_header_11b
                cell.fill = fill_header_gray
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border_cell
            else:
                if cell.value is not None and str(cell.value) != "":
                    if cell.fill.fill_type is None:
                        cell.font = font_data_11r
                    cell.border = border_cell
                    if c in (7, 8, 9, 10, 12, 13, 14, 15):
                        cell.number_format = FMT_NORM
                    if c in (1, 2, 3, 5):
                        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Thiết lập độ rộng cột chuẩn thoáng
    col_widths = {
        'A': 8, 'B': 16, 'C': 14, 'D': 45, 'E': 10, 'F': 14,
        'G': 12, 'H': 12, 'I': 12, 'J': 12, 'K': 10,
        'L': 14, 'M': 14, 'N': 14, 'O': 14, 'P': 28
    }
    for col_letter, width in col_widths.items():
        ws_hp.column_dimensions[col_letter].width = width

    return ws_hp, last_hp_row