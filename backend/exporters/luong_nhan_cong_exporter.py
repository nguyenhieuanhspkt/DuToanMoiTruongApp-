from openpyxl.drawing.image import Image
import os
import openpyxl
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
)
from openpyxl.worksheet.worksheet import Worksheet
from typing import Dict, Tuple

# ==========================================================
# DỮ LIỆU THEO TTLT 56/2015 & THÔNG TƯ 41/2014/TT-BTNMT
# ==========================================================

NGACH_NHAN_CONG = [
    {
        "stt": 1,
        "ma_so": "V.06.05.13",
        "ten": "Quan trắc viên TNMT hạng II",
        "he_so": 4.40,
    },
    {
        "stt": 2,
        "ma_so": "V.06.05.14",
        "ten": "Quan trắc viên TNMT hạng III",
        "he_so": 2.34,
    },
    {
        "stt": 3,
        "ma_so": "V.06.05.15",
        "ten": "Quan trắc viên TNMT hạng IV",
        "he_so": 1.86,
    },
]


# ==========================================================
# SHEET
# ==========================================================

def _init_sheet(wb: openpyxl.Workbook) -> Worksheet:
    """Khởi tạo hoặc làm mới sheet 'Tính lương NC' và bật hiển thị lưới."""
    name = "Tính lương NC"
    if name in wb.sheetnames:
        ws = wb[name]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(name)
    
    ws.views.sheetView[0].showGridLines = True
    return ws


# ==========================================================
# HEADER
# ==========================================================

def _write_header(ws: Worksheet):
    """Ghi tiêu đề dự án và header bảng lương chi tiết (Dòng 1 đến 5)."""
    ws.append(["BẢNG LƯƠNG VÀ ĐƠN GIÁ NGÀY CÔNG NHÂN CÔNG QUAN TRẮC TNMT"])
    ws.append(["='Công trình'!A3"])
    ws.append(["='Công trình'!A4"])
    ws.append([]) # Dòng 4 trống
    ws.append([
        "STT",
        "Mã số",
        "Chức danh",
        "Hệ số lương",
        "Lương cơ sở",
        "Lương tháng (VNĐ)",
        "PC Lưu động (0.4)",
        "PC Độc hại (0.1)",
        "Tổng thu nhập tháng",
        "Số ngày công",
        "Đơn giá ngày công (VNĐ)",
        "Căn cứ pháp lý",
    ])


# ==========================================================
# DATA
# ==========================================================

def _write_data(ws: Worksheet, luong_co_so: float, ngay_cong_chuan: float = 26.0) -> Dict[str, int]:
    """
    Ghi dữ liệu các ngạch nhân công, tích hợp phụ cấp theo TT 41/2014 
    và tính đơn giá ngày công, trả về từ điển ánh xạ tọa độ dòng (code -> row_index).
    """
    start_row = 6
    output_row = start_row
    nc_rows = {}

    for item in NGACH_NHAN_CONG:
        ma_so_key = str(item["ma_so"]).lower().strip()
        nc_rows[ma_so_key] = output_row  # Lưu tọa độ dòng thực tế trên Excel

        ws.append([
            item["stt"],
            item["ma_so"],
            item["ten"],
            item["he_so"],
            luong_co_so,
            f"=ROUND(D{output_row}*E{output_row},0)",          # Cột F: Lương tháng
            f"=ROUND(0.4*E{output_row},0)",                    # Cột G: Phụ cấp lưu động (Hệ số 0.4)
            f"=ROUND(0.1*E{output_row},0)",                    # Cột H: Phụ cấp độc hại (Hệ số 0.1)
            f"=ROUND(F{output_row}+G{output_row}+H{output_row},0)", # Cột I: Tổng thu nhập tháng
            ngay_cong_chuan,                                   # Cột J: Số ngày công chuẩn (26 ngày)
            f"=ROUND(I{output_row}/J{output_row},0)",          # Cột K: Đơn giá ngày công
            "TTLT 56/2015 & TT 41/2014/TT-BTNMT",              # Cột L: Căn cứ pháp lý
        ])
        output_row += 1

    return nc_rows


# ==========================================================
# TOTAL
# ==========================================================

def _write_total(ws: Worksheet) -> int:
    """Ghi dòng tổng cộng giá trị lương và đơn giá."""
    last_row = ws.max_row
    total_row = last_row + 1

    ws.append([
        "",
        "",
        "TỔNG CỘNG / TRUNG BÌNH",
        "",
        "",
        f"=SUM(F6:F{last_row})",
        f"=SUM(G6:G{last_row})",
        f"=SUM(H6:H{last_row})",
        f"=SUM(I6:I{last_row})",
        "",
        f"=ROUND(AVERAGE(K6:K{last_row}),0)",
        "",
    ])
    return total_row


# ==========================================================
# FORMAT
# ==========================================================

def _format_sheet(
    ws: Worksheet,
    font_title_15b: Font,
    font_sub_11i: Font,
    font_header_11b: Font,
    font_data_11r: Font,
    font_data_11b: Font,
    fill_header_gray: PatternFill,
    border_cell: Border,
    border_total: Border,
    total_row: int,
):
    """Định dạng giao diện, phông chữ, màu sắc, viền và độ rộng cột cho sheet Tính lương NC."""
    ws.merge_cells("A1:L1")
    ws["A1"].font = font_title_15b
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A2"].font = font_sub_11i
    ws["A3"].font = font_sub_11i

    # Header bảng (Dòng 5)
    header_row = 5
    ws.row_dimensions[header_row].height = 30
    for c in range(1, 13):
        cell = ws.cell(header_row, c)
        cell.font = font_header_11b
        cell.fill = fill_header_gray
        cell.border = border_cell
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Dữ liệu và Dòng tổng
    for r in range(6, ws.max_row + 1):
        is_total = (r == total_row)
        for c in range(1, 13):
            cell = ws.cell(r, c)
            cell.font = font_data_11b if is_total else font_data_11r
            cell.border = border_total if is_total else border_cell

            if is_total:
                cell.fill = fill_header_gray

            if c in (1, 2, 4, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (5, 6, 7, 8, 9, 11):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    widths = {
        "A": 8,
        "B": 15,
        "C": 35,
        "D": 15,
        "E": 18,
        "F": 20,
        "G": 20,
        "H": 20,
        "I": 22,
        "J": 15,
        "K": 22,
        "G": 30,
        "L": 35,
    }
    for col, width in widths.items():
        if col in ws.column_dimensions:
            ws.column_dimensions[col].width = width

    ws.freeze_panes = "A6"


# ==========================================================
# MAIN BUILDER
# ==========================================================

def build_luong_nhan_cong_sheet(
    wb: openpyxl.Workbook,
    font_title_15b: Font,
    font_sub_11i: Font,
    font_header_11b: Font,
    font_data_11r: Font,
    font_data_11b: Font,
    fill_header_gray: PatternFill,
    border_cell: Border,
    border_total: Border,
    luong_co_so: float = 2_340_000,
    ngay_cong_chuan: float = 26.0,
) -> Tuple[Worksheet, int, Dict[str, int]]:
    """
    Hàm điều phối chính: Khởi tạo, viết header, tính toán lương + phụ cấp + đơn giá ngày công.
    Trả về: (worksheet, total_row, nc_rows) trong đó nc_rows chứa tọa độ trỏ đến Cột K (Đơn giá ngày công).
    """
    ws = _init_sheet(wb)
    _write_header(ws)
    nc_rows = _write_data(ws, luong_co_so, ngay_cong_chuan)
    total_row = _write_total(ws)
    
    _format_sheet(
        ws,
        font_title_15b,
        font_sub_11i,
        font_header_11b,
        font_data_11r,
        font_data_11b,
        fill_header_gray,
        border_cell,
        border_total,
        total_row,
    )
    # ==========================================================
    # CHÈN HÌNH ẢNH GIẢI TRÌNH (Hình 1 & Hình 2)
    # ==========================================================
    try:
        # Đường dẫn tới thư mục chứa ảnh
        img_dir = os.path.join("backend", "data", "luongNC")
        
        img_path_1 = os.path.join(img_dir, "Hình 1.png")
        img_path_2 = os.path.join(img_dir, "Hình 2.png")

        start_row_for_images = total_row + 3  # Cách dòng tổng cộng 3 dòng

        if os.path.exists(img_path_1):
            img1 = Image(img_path_1)
            # Có thể thu nhỏ kích thước ảnh nếu quá lớn (ví dụ: width=600, height=400)
            img1.width = 650
            img1.height = 450
            ws.add_image(img1, f"A{start_row_for_images}")
            
            # Tính dòng tiếp theo cho Hình 2 dựa trên chiều cao ảnh ước lượng
            start_row_for_images += 26 # Khoảng cách dòng tương ứng với chiều cao ảnh

        if os.path.exists(img_path_2):
            img2 = Image(img_path_2)
            img2.width = 650
            img2.height = 450
            ws.add_image(img2, f"A{start_row_for_images}")

    except Exception as e:
        print(f"⚠️ Không thể chèn hình ảnh giải trình vào sheet Tính lương NC: {e}")

    return ws, total_row, nc_rows