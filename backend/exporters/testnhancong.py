import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Import hàm cần test từ module của anh (điều chỉnh lại đường dẫn import cho đúng thực tế dự án)
# Ví dụ: from backend.exporters.nhan_cong_exporter import build_nhan_cong_sheet
try:
    from backend.exporters.nhan_cong_exporter import build_nhan_cong_sheet
except ImportError:
    # Nếu đặt cùng thư mục hoặc cấu trúc khác, anh sửa lại dòng import này cho phù hợp
    from nhan_cong_exporter import build_nhan_cong_sheet


# --- 1. MOCK DTO CLASSES (GIẢ LẬP CẤU TRÚC DỮ LIỆU ĐẦU VÀO) ---
class DetailItem:
    def __init__(self, name, unit, norm):
        self.name = name
        self.unit = unit
        self.norm = norm

class Group:
    def __init__(self, items):
        self.items = items

class JobDTO:
    def __init__(self, stt, code, name, row_index, groups):
        self.stt = stt
        self.code = code
        self.name = name
        self.row_index_in_cong_trinh = row_index
        self.groups = groups


def run_unit_test():
    print("đang khởi tạo Workbook giả lập để test sheet Nhân công...")
    wb = openpyxl.Workbook()

    # --- 2. TẠO SHEET 'Công trình' GIẢ LẬP ĐỂ TEST LINK CÔNG THỨC ---
    ws_ct = wb.active
    ws_ct.title = "Công trình"
    ws_ct.append(["BẢNG CHI TIẾT KHỐI LƯỢNG CÔNG TRÌNH QUAN TRẮC MÔI TRƯỜNG"])
    ws_ct.append(["(Căn cứ Quyết định...)"])
    ws_ct.append(["CÔNG TRÌNH: Quan trắc môi trường cho NMNĐ Vĩnh Tân 4"])
    ws_ct.append(["HẠNG MỤC: Toàn bộ khối lượng quan trắc"])
    ws_ct.append([])
    # Dòng 6: Header sheet Công trình
    ws_ct.append(["STT", "Ký hiệu vị trí", "Mã hiệu công tác", "Danh mục công tác quan trắc", "Đơn vị"])
    # Dòng 7: Dòng tiêu đề phân mục (giả lập để test xem bộ lọc có bỏ qua không)
    ws_ct.append(["", "", "", "PHẦN: GIÁM SÁT KHÔNG KHÍ XUNG QUANH", ""])
    # Dòng 8: Công tác thực tế (STT = 1, row_index = 8)
    ws_ct.append(["1", "K1", "KK1a", "Nhiệt độ", "°C"])
    # Dòng 9: Công tác thực tế (STT = 2, row_index = 9)
    ws_ct.append(["2", "K1", "KK1b", "Độ ẩm", "%"])

    # --- 3. TẠO SHEET 'Hao phí VT' GIẢ LẬP ---
    ws_hp = wb.create_sheet(title="Hao phí VT")
    ws_hp.append(["BẢNG HAO PHÍ VẬT TƯ"])
    ws_hp.append([])
    ws_hp.append([])
    ws_hp.append([])
    ws_hp.append([])
    # Header dòng 6
    ws_hp.append(["STT", "Mã hiệu", "Tên vật tư", "Đơn vị", "Tên nhân công/máy", "", "", "", "Khối lượng"])
    # Dòng dữ liệu mẫu 6
    ws_hp.cell(row=6, column=2, value="KK1a")
    ws_hp.cell(row=6, column=9, value=1.5)

    last_hp_row = 6

    # --- 4. TẠO DỮ LIỆU JOBS_DTOS MẪU ---
    # Tạo detail item nhân công khớp với tên sẽ bóc tách
    nc_item = DetailItem(name="Nhân công khảo sát bậc 3/7", unit="công", norm=0.5)
    
    job_1 = JobDTO(
        stt="1",
        code="KK1a",
        name="Nhiệt độ",
        row_index=8, # Trỏ tương ứng dòng 8 bên sheet Công trình
        groups={"NC": Group([nc_item])}
    )

    jobs_dtos = [job_1]

    # --- 5. KHỞI TẠO CÁC STYLE TIÊU CHUẨN MỞ RỘNG ---
    font_title_15b = Font(name="Times New Roman", size=15, bold=True)
    font_sub_11i = Font(name="Times New Roman", size=11, italic=True)
    font_header_11b = Font(name="Times New Roman", size=11, bold=True)
    font_data_11r = Font(name="Times New Roman", size=11, bold=False)
    font_data_11b = Font(name="Times New Roman", size=11, bold=True)

    fill_header_gray = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    fill_sub_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_total_pale = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_border = Side(border_style="thin", color="000000")
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_total = Border(top=thin_border, bottom=Side(border_style="double", color="000000"))

    FMT_CURRENCY = "#,##0"
    FMT_NORM = "#,##0.0000"

    print("Đang gọi hàm build_nhan_cong_sheet...")
    
    # --- 6. THỰC THI HÀM CẦN TEST ---
    try:
        ws_nc, total_row = build_nhan_cong_sheet(
            wb=wb,
            last_hp_row=last_hp_row,
            jobs_dtos=jobs_dtos,
            font_title_15b=font_title_15b,
            font_sub_11i=font_sub_11i,
            font_header_11b=font_header_11b,
            font_data_11r=font_data_11r,
            font_data_11b=font_data_11b,
            fill_header_gray=fill_header_gray,
            fill_sub_gray=fill_sub_gray,
            fill_total_pale=fill_total_pale,
            border_cell=border_cell,
            border_total=border_total,
            FMT_CURRENCY=FMT_CURRENCY,
            FMT_NORM=FMT_NORM
        )
        
        output_filename = "test_nhan_cong_output.xlsx"
        wb.save(output_filename)
        print(f"Test THÀNH CÔNG! File kết quả đã được lưu tại: {output_filename}")
        print(f"Dòng tổng kết nằm ở hàng số: {total_row}")

    except Exception as e:
        print(f"Test THẤT BẠI do lỗi exception: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    run_unit_test()